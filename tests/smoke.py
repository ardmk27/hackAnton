#!/usr/bin/env python3
"""
Uji asap HTTP end-to-end (hanya pustaka standar Python).
  php -S 127.0.0.1:8080 -t .   (di folder proyek, dengan database terisi dari schema.sql)
  python3 tests/smoke.py http://127.0.0.1:8080
Mencakup: daftar, masuk, data demo, periksa+koreksi, konfirmasi, tren, ekspor, buka kembali sesi, isolasi antar Posyandu.
"""
import html, re, sys, os, subprocess, base64, uuid, urllib.parse, urllib.request, http.cookiejar, zipfile, io
from html.parser import HTMLParser

BASE = (sys.argv[1] if len(sys.argv) > 1 else 'http://127.0.0.1:8080').rstrip('/')
ok = fail = 0


def check(label, cond, detail=''):
    global ok, fail
    if cond:
        ok += 1
        print('  ok   ' + label)
    else:
        fail += 1
        print('  GAGAL ' + label + ('  -> ' + detail if detail else ''))


class NoRedirect(urllib.request.HTTPRedirectHandler):
    def redirect_request(self, *a, **k):
        return None


class Client:
    def __init__(self):
        self.jar = http.cookiejar.CookieJar()
        self.op = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(self.jar))
        self.op_nr = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(self.jar), NoRedirect)

    def get(self, path, follow=True, raw=False):
        op = self.op if follow else self.op_nr
        try:
            r = op.open(BASE + '/' + path.lstrip('/'))
            data = r.read()
            return r.status, (data if raw else data.decode('utf-8', 'replace')), r.headers, r.geturl()
        except urllib.error.HTTPError as e:
            data = e.read()
            return e.code, (data if raw else data.decode('utf-8', 'replace')), e.headers, e.geturl()

    def post(self, path, fields, follow=True):
        body = urllib.parse.urlencode(fields, doseq=True).encode()
        op = self.op if follow else self.op_nr
        req = urllib.request.Request(BASE + '/' + path.lstrip('/'), data=body)
        try:
            r = op.open(req)
            return r.status, r.read().decode('utf-8', 'replace'), r.headers, r.geturl()
        except urllib.error.HTTPError as e:
            return e.code, e.read().decode('utf-8', 'replace'), e.headers, e.geturl()


def multipart(fields, files):
    b = uuid.uuid4().hex
    parts = []
    for k, v in fields.items():
        parts.append(('--%s\r\nContent-Disposition: form-data; name="%s"\r\n\r\n%s\r\n' % (b, k, v)).encode())
    for k, (fn, data, ct) in files.items():
        parts.append(('--%s\r\nContent-Disposition: form-data; name="%s"; filename="%s"\r\nContent-Type: %s\r\n\r\n' % (b, k, fn, ct)).encode() + data + b'\r\n')
    parts.append(('--%s--\r\n' % b).encode())
    return b''.join(parts), 'multipart/form-data; boundary=' + b


def post_multipart(client, path, fields, files):
    body, ctype = multipart(fields, files)
    req = urllib.request.Request(BASE + '/' + path.lstrip('/'), data=body, headers={'Content-Type': ctype})
    try:
        r = client.op.open(req)
        return r.status, r.read().decode('utf-8', 'replace'), r.headers, r.geturl()
    except urllib.error.HTTPError as e:
        return e.code, e.read().decode('utf-8', 'replace'), e.headers, e.geturl()


# PNG 1x1 valid
TINY_PNG = base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg==')


class FormParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.fields = []   # (name, value)
        self._sel = None
        self._sel_val = None

    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        if tag == 'input':
            name = a.get('name')
            t = (a.get('type') or 'text').lower()
            if not name or t in ('submit', 'button'):
                return
            if t == 'checkbox':
                if 'checked' in a:
                    self.fields.append((name, a.get('value', '1')))
                return
            self.fields.append((name, a.get('value', '') or ''))
        elif tag == 'select' and a.get('name'):
            self._sel = a['name']
            self._sel_val = ''
        elif tag == 'option' and self._sel is not None and 'selected' in a:
            self._sel_val = a.get('value', '')

    def handle_endtag(self, tag):
        if tag == 'select' and self._sel is not None:
            self.fields.append((self._sel, self._sel_val))
            self._sel = None


def csrf(page):
    m = re.search(r'name="csrf" value="([0-9a-f]+)"', page)
    return m.group(1) if m else ''


def form_fields(page, form_id):
    m = re.search(r'<form[^>]*id="%s".*?</form>' % form_id, page, re.S)
    p = FormParser()
    p.feed(m.group(0) if m else '')
    return p.fields


def register_and_login(c, name):
    s, pg, _, _ = c.get('register.php')
    s, pg, _, _ = c.post('register.php', {'csrf': csrf(pg), 'name': name})
    code = re.search(r'(PSY-[0-9A-F]{8})', pg)
    access = re.search(r'([A-Z2-9]{4}-[A-Z2-9]{4}-[A-Z2-9]{4})', pg)
    check('registrasi menampilkan ID dan kode akses', bool(code and access))
    s, pg, _, _ = c.get('login.php')
    s, pg, _, url = c.post('login.php', {'csrf': csrf(pg), 'posyandu_code': code.group(1), 'access_code': access.group(1)})
    check('login berhasil ke dashboard', 'dashboard.php' in url, url)
    return code.group(1), access.group(1)


print('== Akses ==')
c = Client()
s, pg, _, url = c.get('dashboard.php')
check('tanpa login diarahkan ke login', 'login.php' in url)
s, _, _, _ = c.get('anak_detail.php?id=1')
check('halaman anak tanpa login tidak terbuka', 'login.php' in c.get('anak_detail.php?id=1')[3])
code, access = register_and_login(c, 'Posyandu Uji Asap')

print('== Beranda & data demo ==')
s, pg, _, _ = c.get('dashboard.php')
check('beranda memuat tombol data demo', 'Muat data demo' in pg)
s, pg2, _, url = c.post('demo_seed.php', {'csrf': csrf(pg), 'mode': 'story'})
check('data demo mengarahkan ke sesi draf', 'review.php?id=' in url, url)
sesi_id = int(re.search(r'id=(\d+)', url).group(1))
check('review memuat saran nama untuk "M. Rizky"', 'Mungkin anak ini' in pg2 and 'Muhammad Rizki' in pg2)
check('review menandai BB Budi yang janggal dari selisih catatan sebelumnya', 'BB naik 6,2 kg' in pg2)
check('review menampilkan konteks catatan sebelumnya', 'Catatan sebelumnya' in pg2)
check('review memuat foto sumber hanya jika ada (demo tanpa foto)', 'photo.php' not in pg2)

print('== Konfirmasi dengan koreksi ==')
fields = form_fields(pg2, 'review-form')
names = [n for n, _ in fields]
check('form berisi 14 baris', sum(1 for n in names if n.endswith('[nama_peserta]')) == 14)
new = []
row_of = {}
for n, v in fields:
    m = re.match(r'rows\[(\w+)\]\[(\w+)\]', n)
    if m:
        row_of.setdefault(m.group(1), {})[m.group(2)] = v
rizki_key = next(k for k, r in row_of.items() if r.get('nama_peserta') == 'M. Rizky')
budi_key = next(k for k, r in row_of.items() if r.get('nama_peserta') == 'Budi Santoso')
for n, v in fields:
    m = re.match(r'rows\[(\w+)\]\[(\w+)\]', n)
    if m:
        k, f = m.groups()
        if k == rizki_key:
            v = {'nama_peserta': 'Muhammad Rizki', 'nomor_pendaftaran': '003', 'tanggal_lahir': '2025-03-12', 'jenis_kelamin': 'L', 'tb_pb_cm': '78', 'human_checked': '1'}.get(f, v)
        if k == budi_key:
            v = {'bb_kg': '11,4', 'human_checked': '1'}.get(f, v)
    new.append((n, v))
new = [(n, v) for n, v in new if not n.endswith('[needs_review]')]
new += [('csrf', csrf(pg2)), ('sesi_id', str(sesi_id)), ('action', 'confirm')]
s, pg3, _, url = c.post('confirm.php', new)
check('konfirmasi sukses ke detail sesi', 'session_detail.php' in url and 'dikonfirmasi' in pg3, pg3[:200])
check('pesan menyebut jumlah baris yang dikoreksi', '2 baris Anda koreksi' in pg3, re.sub(r'<[^>]+>', ' ', pg3)[:300])

print('== Tren & perhatian ==')
s, pg, _, _ = c.get('attention.php')
check('Rizki muncul sebagai Prioritas', 'Muhammad Rizki' in pg and 'Prioritas' in pg and 'tidak naik 2 bulan berturut-turut' in pg)
check('Dewi muncul sebagai Cek', 'Dewi Lestari' in pg, re.sub(r'\s+', ' ', re.sub(r'<[^>]+>', ' ', pg))[300:900])
check('anak yang naik tidak muncul', 'Siti Aminah' not in pg)
txt = re.sub(r'\s+', ' ', re.sub(r'<[^>]+>', ' ', pg))
check('Rina (tanpa catatan Oktober) tidak disimpulkan absen', 'Rina Wulandari' not in pg and 'berarti mereka tidak hadir' in txt)
s, pg, _, _ = c.get('attention.php?level=prioritas')
check('filter Prioritas hanya Rizki', 'Muhammad Rizki' in pg and 'Dewi Lestari' not in pg)
rid = re.search(r'anak_detail\.php\?id=(\d+)&amp;periode=2026-10"[^>]*>Muhammad Rizki', pg)
rid = rid.group(1) if rid else None
check('tautan ke detail anak ada', rid is not None)
s, pg, _, _ = c.get('anak_detail.php?id=%s&periode=2026-10' % rid)
check('detail anak memuat grafik SVG', '<svg class="kcf-chart"' in pg and 'tidak naik' in pg)
check('detail anak memuat draf pesan tanpa diagnosis', 'Draf pesan untuk orang tua' in pg and 'stunting' not in pg.lower())
check('detail anak jujur bahwa status gizi belum dihitung', 'Belum dihitung' in pg)
s, pg, _, _ = c.get('anak.php?q=dewi')
check('pencarian anak bekerja', 'Dewi Lestari' in pg and 'Siti Aminah' not in pg)
s, pg, _, _ = c.get('dashboard.php')
check('beranda memuat panel perhatian', 'Perlu perhatian bulan ini' in pg and 'Muhammad Rizki' in pg)
check('beranda menampilkan angka koreksi', 'perlu dikoreksi kader' in pg)

print('== Isolasi antar Posyandu ==')
c2 = Client()
register_and_login(c2, 'Posyandu Lain')
s, pg, _, _ = c2.get('anak_detail.php?id=%s' % rid)
check('Posyandu lain tidak bisa membuka anak milik orang lain', s == 404, str(s))
s, pg, _, _ = c2.get('review.php?id=%d' % sesi_id)
check('Posyandu lain tidak bisa membuka sesi orang lain', s == 404, str(s))
s, pg, _, _ = c2.get('attention.php')
check('daftar perhatian Posyandu lain kosong', 'Muhammad Rizki' not in pg)


print('== Rekap & ekspor ==')
c = c_main = c  # klien Posyandu utama
s, pg, _, _ = c.get('recap.php')
check('rekap per periode menampilkan 3 bulan', all(x in pg for x in ['Oktober 2026', 'September 2026', 'Agustus 2026']))
check('rekap menampilkan penanda prioritas dan cek', '1 prioritas' in pg and '1 cek' in pg)
check('rekap menampilkan tab jenis formulir', 'Balita (' in pg and 'Ibu hamil (' in pg and 'Lansia (' in pg)
s, pg, _, _ = c.get('recap.php?periode=2026-10')
check('detail rekap menawarkan PDF dan XLSX', 'export_pdf.php?periode=2026-10' in pg and 'export_xlsx.php?periode=2026-10' in pg)
check('CSV tidak ditawarkan saat XLSX tersedia', 'export_csv.php' not in pg)
check('detail rekap memuat selisih BB dan penanda', 'Prioritas' in pg and 'BB tidak naik' in pg)

s, body, hdr, _ = c.get('export_pdf.php?periode=2026-10', raw=True)
check('PDF terunduh dengan tipe yang benar', s == 200 and 'application/pdf' in (hdr.get('Content-Type') or '') and body[:5] == b'%PDF-', str(s))
try:
    from pypdf import PdfReader
    rd = PdfReader(io.BytesIO(body))
    txt = ' '.join(pg_.extract_text() for pg_ in rd.pages)
    check('PDF dapat dibaca dan memuat judul, data, dan penanda', 'Laporan Penimbangan dan Pengukuran Balita' in txt and 'Muhammad Rizki' in txt and 'Prioritas' in txt)
    check('PDF memuat kolom tanda tangan dan penafian', 'Bidan / Petugas Puskesmas' in txt and 'bukan diagnosis' in re.sub(r'\s+', ' ', txt))
    check('PDF memuat nomor halaman', re.search(r'Halaman 1 dari \d+', txt) is not None)
except ImportError:
    print('  (lewati: pypdf tidak terpasang)')

s, body, hdr, _ = c.get('export_xlsx.php?periode=2026-10', raw=True)
check('XLSX terunduh dengan tipe yang benar', s == 200 and 'spreadsheetml' in (hdr.get('Content-Type') or ''), str(s))
try:
    import openpyxl, datetime
    wb = openpyxl.load_workbook(io.BytesIO(body))
    check('XLSX valid: sheet Laporan, Ringkasan, Data', wb.sheetnames == ['Laporan', 'Ringkasan', 'Data'], str(wb.sheetnames))
    ws = wb['Laporan']
    check('Laporan: judul dan subjudul tergabung', 'Laporan Penimbangan' in str(ws['A1'].value) and 'PSY-' in str(ws['A2'].value))
    hdrs = [c_.value for c_ in ws[4]]
    check('Laporan: header kolom pada baris 4', hdrs[:2] == ['No.', 'Nama'] and 'Penanda' in hdrs, str(hdrs))
    rows = {r[1].value: r for r in ws.iter_rows(min_row=5) if r[1].value}
    rz = rows.get('Muhammad Rizki')
    bb = hdrs.index('BB (kg)')
    check('Laporan: BB tersimpan sebagai ANGKA', rz is not None and isinstance(rz[bb].value, (int, float)) and abs(rz[bb].value - 9.0) < 1e-9)
    check('Laporan: tanggal lahir adalah tanggal sungguhan', isinstance(rz[hdrs.index('Tgl lahir')].value, datetime.datetime))
    check('Laporan: penanda Prioritas dan selisih 0', rz[hdrs.index('Penanda')].value == 'Prioritas' and rz[hdrs.index('Selisih BB (kg)')].value == 0)
    check('Laporan: header membeku, ada filter, dan siap cetak lanskap', ws.freeze_panes == 'A5' and ws.auto_filter.ref and ws.page_setup.orientation == 'landscape' and ws.print_title_rows == '$4:$4')
    check('Laporan: baris selang-seling dan gaya header', ws['A4'].fill.fgColor.rgb.endswith('1E2D87') and ws['A5'].fill.fgColor.rgb != ws['A6'].fill.fgColor.rgb)
    alltxt = ' '.join(str(c_.value) for r in ws.iter_rows() for c_ in r if c_.value is not None)
    check('Laporan: catatan menyatakan BUKAN diagnosis', 'BUKAN diagnosis' in alltxt)
    sm = wb['Ringkasan']
    txt = ' '.join(str(c_.value) for r in sm.iter_rows() for c_ in r if c_.value is not None)
    check('Ringkasan memuat angka utama dan informasi', 'Catatan terkonfirmasi' in txt and 'Posyandu Uji Asap' in txt and 'Dibuat pada' in txt)
    dt = wb['Data']
    dh = [c_.value for c_ in dt[1]]
    check('Data: kolom mesin lengkap', dh[:4] == ['posyandu_id', 'sesi_id', 'nomor_pendaftaran', 'nama_peserta'] and 'bb_kg' in dh and 'penanda_untuk_ditinjau' in dh)
except ImportError:
    print('  (lewati: openpyxl tidak terpasang)')

print('== Buka kembali sesi & jejak ==')
s, pg, _, _ = c.get('session_detail.php?id=%d' % sesi_id)
check('detail sesi terkonfirmasi menawarkan "Buka kembali"', 'Buka kembali' in pg or 'buka kembali' in pg.lower())
s, pg2, _, url = c.post('session_detail.php?id=%d' % sesi_id, {'csrf': csrf(pg), 'action': 'reopen', 'reason': ''})
check('buka kembali tanpa alasan ditolak', 'alasan' in pg2.lower() and 'sesi_foto' not in pg2.lower())
s, pg2, _, url = c.post('session_detail.php?id=%d' % sesi_id, {'csrf': csrf(pg), 'action': 'reopen', 'reason': 'Salah baca BB anak nomor 5'})
check('buka kembali dengan alasan → halaman periksa', 'review.php?id=%d' % sesi_id in url, url)
s, pg, _, _ = c.get('attention.php')
check('setelah dibuka kembali, data Oktober keluar dari tren (kembali ke draf)', 'Muhammad Rizki' not in pg or 'tidak naik 2 bulan' not in pg)
s, pg, _, _ = c.get('recap.php')
check('rekap Oktober hilang sampai dikonfirmasi ulang', 'Oktober 2026' not in pg)
s, pg, _, _ = c.get('review.php?id=%d' % sesi_id)
check('banner menyebut sesi dibuka kembali', 'dibuka kembali untuk revisi' in pg)
fields = form_fields(pg, 'review-form')
fields = [(n, v) for n, v in fields if not n.endswith('[needs_review]')]
fields += [('csrf', csrf(pg)), ('sesi_id', str(sesi_id)), ('action', 'confirm')]
s, pg, _, url = c.post('confirm.php', fields)
check('konfirmasi ulang berhasil', 'session_detail.php' in url and 'dikonfirmasi' in pg)
s, pg, _, _ = c.get('attention.php')
check('tren kembali muncul setelah konfirmasi ulang', 'Muhammad Rizki' in pg)
s, pg, _, _ = c.get('session_detail.php?id=%d' % sesi_id)
check('detail sesi menampilkan jejak aktivitas', 'Jejak aktivitas' in pg and 'dibuka kembali' in pg.lower())

print('== Keamanan ekspor ==')
c2 = c2
for path in ['export_pdf.php?periode=2026-10', 'export_xlsx.php?periode=2026-10', 'export_csv.php?periode=2026-10']:
    s, body, _, _ = c2.get(path, raw=True)
    check('Posyandu lain tidak melihat data lewat ' + path.split('?')[0], b'Rizki' not in body and b'Muhammad' not in body)
s, _, _, _ = c2.get('export_xlsx.php?sesi_id=%d' % sesi_id)
check('ekspor XLSX sesi milik orang lain → 404', s == 404, str(s))
s, _, _, _ = c2.get('export_pdf.php?sesi_id=%d' % sesi_id)
check('ekspor PDF sesi milik orang lain → 404', s == 404, str(s))
c3 = Client()
for path in ['export_pdf.php?periode=2026-10', 'export_xlsx.php?periode=2026-10']:
    s, _, _, url = c3.get(path)
    check('tanpa login diarahkan ke login: ' + path.split('?')[0], 'login.php' in url)

print('== Nama kembar ==')
fx = subprocess.run(['php', os.path.join(os.path.dirname(os.path.abspath(__file__)), 'fixtures', 'make_dupes.php'), code], capture_output=True, text=True)
check('fixture nama kembar dibuat', fx.stdout.strip().isdigit(), fx.stderr[:200])
dsid = int(fx.stdout.strip())
s, pg, _, _ = c.get('review.php?id=%d' % dsid)
check('review menampilkan pemilih untuk nama kembar', 'Nama kembar, pilih yang benar' in pg)
check('pemilih memuat pembeda kedua anak', 'no. 001' in pg and 'no. 016' in pg and 'lahir 14/02/2025' in pg and 'lahir 09/09/2025' in pg)
check('peringatan jelas pada baris nama kembar', 'Ada 2 anak bernama Siti Aminah' in pg)
fields = form_fields(pg, 'review-form')
conf = [(n, v) for n, v in fields if not n.endswith('[needs_review]')] + [('csrf', csrf(pg)), ('sesi_id', str(dsid)), ('action', 'confirm')]
s, pg2, _, url = c.post('confirm.php', conf)
txt = re.sub(r'\s+', ' ', re.sub(r'<[^>]+>', ' ', pg2))
check('konfirmasi DIBLOKIR saat nama kembar tidak dibedakan', 'review.php' in url and 'Nama sama dengan 2 anak' in txt, url)
check('baris kembar dalam halaman ikut ditandai', 'Nama sama dengan baris' in txt)
s, pg3, _, _ = c.get('session_detail.php?id=%d' % dsid)
check('sesi tetap draf (tidak ada yang tersimpan)', 'Draf' in pg3 and 'Terkonfirmasi' not in re.sub(r'<[^>]+>', ' ', pg3).split('Detail sesi')[1][:200])
# perbaiki: pilih Siti no. 016; dua Rudi dibedakan dengan no. daftar (anak kembar sungguhan)
rowkeys = {}
for n, v in fields:
    m = re.match(r'rows\[(\w+)\]\[(\w+)\]', n)
    if m:
        rowkeys.setdefault(m.group(1), {})[m.group(2)] = v
ks = [k for k, r in rowkeys.items()]
siti = next(k for k in ks if rowkeys[k].get('nama_peserta') == 'Siti Aminah')
rudis = [k for k in ks if rowkeys[k].get('nama_peserta') == 'Rudi Hartono']
fix = []
seen = 0
for n, v in fields:
    m = re.match(r'rows\[(\w+)\]\[(\w+)\]', n)
    if m and not n.endswith('[needs_review]'):
        k, f = m.groups()
        if k == siti and f == 'nomor_pendaftaran': v = '016'
        if k == siti and f == 'tanggal_lahir': v = '2025-09-09'
        if k in rudis and f == 'nomor_pendaftaran': v = 'R%d' % (rudis.index(k) + 1)
        fix.append((n, v))
    elif not m:
        fix.append((n, v))
fix = [(n, v) for n, v in fix if not n.endswith('[needs_review]')] + [('csrf', csrf(pg)), ('sesi_id', str(dsid)), ('action', 'confirm')]
s, pg4, _, url = c.post('confirm.php', fix)
check('setelah dibedakan, konfirmasi berhasil (anak kembar dengan no. beda lolos)', 'session_detail.php' in url and 'dikonfirmasi' in pg4, re.sub(r'<[^>]+>', ' ', pg4)[:300])
s, pg, _, _ = c.get('anak.php?q=siti')
check('daftar anak membedakan dua Siti Aminah', pg.count('Siti Aminah') >= 2 and 'nama sama dengan anak lain' in pg and 'no. 016' in pg)
s, pg, _, _ = c.get('recap.php?periode=2026-11')
check('rekap November memuat tiga baris baru', 'Rudi Hartono' in pg and pg.count('Rudi Hartono') >= 2 and 'Siti Aminah' in pg)

print('== Jenis formulir lain ==')
s, pg, _, _ = c.get('upload.php')
check('unggah menawarkan tiga jenis formulir sebagai pilihan', pg.count('name="form_type"') == 3 and 'Register ibu hamil' in pg and 'Register Posyandu lansia' in pg)
check('formulir contoh ditandai', pg.count('Contoh kolom') >= 2)
s, pg2, _, url = post_multipart(c, 'upload.php', {'csrf': csrf(pg), 'form_type': 'register_ibu_hamil', 'mode': 'mock', 'activity_date': '', 'notes': ''}, {'photo': ('uji.png', TINY_PNG, 'image/png')})
check('unggah formulir ibu hamil (mode demo) menuju halaman periksa', 'review.php?id=' in url, url)
ibu_sid = int(re.search(r'id=(\d+)', url).group(1))
check('review ibu hamil memakai kolom ibu hamil', all(x in pg2 for x in ['Usia kehamilan (minggu)', 'Sistolik (mmHg)', 'Diastolik (mmHg)', 'LiLA (cm)']))
check('review ibu hamil tidak menampilkan kolom balita dan jenis kelamin', 'Lingkar kepala' not in pg2 and 'Jenis kelamin' not in pg2)
check('review ibu hamil menandai kolom sebagai contoh', 'Kolom contoh' in pg2)
check('angka mustahil dikosongkan dan ditandai (TD 1100, usia 120 minggu)', 'Sistolik tidak valid' in pg2 and 'Usia kehamilan tidak valid' in pg2)
fields = form_fields(pg2, 'review-form')
conf = [(n, v) for n, v in fields if not n.endswith('[needs_review]')] + [('csrf', csrf(pg2)), ('sesi_id', str(ibu_sid)), ('action', 'confirm')]
s, pg3, _, url = c.post('confirm.php', conf)
check('konfirmasi register ibu hamil berhasil', 'session_detail.php' in url and 'dikonfirmasi' in pg3, re.sub(r'<[^>]+>', ' ', pg3)[:300])
s, pg, _, _ = c.get('recap.php?form=register_ibu_hamil')
check('rekap ibu hamil terpisah dan memuat periode', 'Oktober 2026' in pg and 'Ibu hamil' in pg)
s, pg, _, _ = c.get('recap.php?periode=2026-10&form=register_ibu_hamil')
check('detail rekap ibu hamil memakai kolom ibu hamil dan tanpa penanda tren', 'Usia kehamilan' in pg and 'Contoh Ibu A' in pg and 'BB tidak naik' not in pg and 'Penanda' not in pg)
s, pg, _, _ = c.get('recap.php?periode=2026-10')
check('rekap balita tidak tercampur data ibu hamil', 'Contoh Ibu A' not in pg)
s, pg, _, _ = c.get('anak.php')
check('daftar anak tidak memuat ibu hamil', 'Contoh Ibu' not in pg)
s, pg, _, _ = c.get('attention.php?periode=2026-10')
check('daftar perhatian tidak terpengaruh formulir lain', 'Muhammad Rizki' in pg and 'Contoh Ibu' not in pg)
s, body, hdr, _ = c.get('export_pdf.php?periode=2026-10&form=register_ibu_hamil', raw=True)
check('PDF ibu hamil terunduh', s == 200 and body[:5] == b'%PDF-')
try:
    txt = re.sub(r'\s+', ' ', ' '.join(pg_.extract_text() for pg_ in PdfReader(io.BytesIO(body)).pages))
    check('PDF ibu hamil memuat judul, kolom, dan catatan contoh', 'Laporan Register ibu hamil' in txt and 'Usia kehamilan' in txt and 'contoh awal' in txt)
except NameError:
    pass
s, body, hdr, _ = c.get('export_xlsx.php?periode=2026-10&form=register_ibu_hamil', raw=True)
try:
    wb = openpyxl.load_workbook(io.BytesIO(body))
    hd = [c_.value for c_ in wb['Laporan'][4]]
    check('XLSX ibu hamil: kolom ibu hamil, tanpa kolom tren', 'Usia kehamilan (minggu)' in hd and 'Sistolik (mmHg)' in hd and 'Penanda' not in hd, str(hd))
    check('XLSX ibu hamil: tekanan darah tersimpan sebagai bilangan bulat', any(r[hd.index('Sistolik (mmHg)')].value == 110 for r in wb['Laporan'].iter_rows(min_row=5) if r[1].value))
except NameError:
    pass

print('\n----------------------------------------')
print('Lulus: %d   Gagal: %d' % (ok, fail))
sys.exit(1 if fail else 0)
